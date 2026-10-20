"""Pure-Python helpers (no Django imports): reading uploads, parsing amounts, duplicate rules."""
import csv
import io
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

MAX_ROWS = 6000
MAX_COLS = 22
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
CENT = Decimal("0.01")
MAX_AMOUNT = Decimal("9999999999.99")  # DecimalField(max_digits=12, decimal_places=2)

# Header names we look for (compared after lower-casing and removing non-alphanumerics).
HEADER_ALIASES = {
    "amount": ("amount", "monto", "importe"),
    "trxn_code": ("trxncode", "transactioncode", "trxcode"),
    "merchant": ("merchants", "merchant", "comercio"),
    "respond_code": ("respondcode", "responsecode", "respcode"),
}
# Fallback positions (0-based) when a header is missing: E, F, G, I
DEFAULT_POSITIONS = {"amount": 4, "trxn_code": 5, "merchant": 6, "respond_code": 8}
FIELD_LABELS = {
    "amount": "Amount",
    "trxn_code": "Trxn Code",
    "merchant": "Merchants",
    "respond_code": "Respond Code",
}
TEXT_LIMITS = {"col1": 200, "col2": 200, "col3": 200, "trxn_code": 50, "merchant": 200, "respond_code": 20}


class UploadError(Exception):
    """An error whose message is safe to show to the user."""


# ---------------------------------------------------------------- cell cleaning

def clean_text(value, limit=None):
    """Cell -> trimmed string. Empty cells become '' (never the text 'nan')."""
    if value is None:
        return ""
    if isinstance(value, float):
        if value != value:  # NaN
            return ""
        if value.is_integer():
            value = int(value)  # 123.0 -> "123"
    elif isinstance(value, datetime):
        value = value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    elif isinstance(value, date):
        value = value.isoformat()
    text = str(value).strip()
    return text[:limit] if limit else text


def _parse_amount_str(s):
    s = s.strip()
    if s.lower() in {"", "nan", "none", "null", "-"}:
        return None
    negative = (s.startswith("(") and s.endswith(")")) or s.startswith("-") or s.endswith("-")
    s = re.sub(r"[^\d.,]", "", s)
    if not re.search(r"\d", s):
        return None
    last_dot, last_comma = s.rfind("."), s.rfind(",")
    dec_sep = None
    if last_dot != -1 and last_comma != -1:
        dec_sep = "." if last_dot > last_comma else ","  # the last separator is the decimal one
    elif last_dot != -1 or last_comma != -1:
        sep = "." if last_dot != -1 else ","
        int_part, _, frac = s.rpartition(sep)
        # one separator: decimal unless it looks like thousands ("1,234"); "0.123" / ".123" stay decimal
        if s.count(sep) == 1 and (len(frac) != 3 or int_part in ("", "0")):
            dec_sep = sep
    if dec_sep:
        int_part, _, frac = s.rpartition(dec_sep)
        s = re.sub(r"[.,]", "", int_part) + "." + frac
    else:
        s = re.sub(r"[.,]", "", s)
    try:
        value = Decimal(s)
    except InvalidOperation:
        return None
    return -value if negative else value


def parse_amount(raw):
    """Return a Decimal rounded to cents, or None if the value can't be understood."""
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, Decimal):
        value = raw
    elif isinstance(raw, int):
        value = Decimal(raw)
    elif isinstance(raw, float):
        if raw != raw or raw in (float("inf"), float("-inf")):
            return None
        value = Decimal(repr(raw))
    else:
        value = _parse_amount_str(str(raw))
        if value is None:
            return None
    try:
        value = value.quantize(CENT, rounding=ROUND_HALF_UP)
    except InvalidOperation:
        return None
    return value if abs(value) <= MAX_AMOUNT else None


# ---------------------------------------------------------------- reading files

def read_table(uploaded):
    """Return the first rows of a CSV/XLSX upload as lists of cells (header row included)."""
    name = (uploaded.name or "").lower()
    if uploaded.size > MAX_UPLOAD_BYTES:
        raise UploadError(f"El archivo supera el máximo de {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
    if name.endswith(".csv"):
        return _read_csv(uploaded)
    if name.endswith(".xlsx"):
        return _read_xlsx(uploaded)
    raise UploadError("Formato no soportado. Usa .csv o .xlsx (guarda los .xls como .xlsx).")


def _read_csv(uploaded):
    raw = uploaded.read()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    first_line = text.split("\n", 1)[0]
    delimiter = max(",;\t|", key=first_line.count)  # Spanish-locale Excel exports use ';'
    if first_line.count(delimiter) == 0:
        delimiter = ","
    rows = []
    try:
        for row in csv.reader(io.StringIO(text, newline=""), delimiter=delimiter):
            rows.append(row)
            if len(rows) > MAX_ROWS + 1:  # header + MAX_ROWS + 1 extra row to detect truncation
                break
    except csv.Error as exc:
        raise UploadError(f"No se pudo leer el CSV: {exc}") from exc
    return rows


def _read_xlsx(uploaded):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
    except Exception as exc:  # BadZipFile, InvalidFileException, ...
        raise UploadError("No se pudo leer el archivo Excel. ¿Está dañado o protegido con contraseña?") from exc
    try:
        rows = []
        for row in workbook.active.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_ROWS + 1:
                break
        return rows
    finally:
        workbook.close()


# ---------------------------------------------------------------- parsing rows

def _norm(header):
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


def detect_columns(header, defaulted=None):
    """Map our fields to column positions using the header names, falling back to E/F/G/I."""
    by_name = {}
    for position, name in reversed(list(enumerate(header))):  # first occurrence wins
        by_name[_norm(name)] = position
    positions = {}
    for field, aliases in HEADER_ALIASES.items():
        position = next((by_name[a] for a in aliases if a in by_name), None)
        if position is None:
            position = DEFAULT_POSITIONS[field]
            if defaulted is not None:
                defaulted.append(f"{FIELD_LABELS[field]} (columna {chr(65 + position)})")
        positions[field] = position
    return positions


def parse_rows(rows):
    """Turn raw rows into Claim field dicts. Returns (records, stats)."""
    stats = {"bad_amounts": 0, "truncated_rows": False, "truncated_cols": False, "defaulted": []}
    if not rows:
        return [], stats
    header, data = list(rows[0]), rows[1:]
    if len(data) > MAX_ROWS:
        data, stats["truncated_rows"] = data[:MAX_ROWS], True
    if len(header) > MAX_COLS or any(len(r) > MAX_COLS for r in data):
        stats["truncated_cols"] = True
    positions = detect_columns(header[:MAX_COLS], stats["defaulted"])

    records = []
    for row in data:
        cells = list(row)[:MAX_COLS]
        if not any(clean_text(c) for c in cells):
            continue  # skip blank lines

        def at(position):
            return cells[position] if 0 <= position < len(cells) else None

        amount_cell = at(positions["amount"])
        amount = parse_amount(amount_cell)
        if amount is None and clean_text(amount_cell):
            stats["bad_amounts"] += 1
        records.append({
            "col1": clean_text(at(0), TEXT_LIMITS["col1"]),
            "col2": clean_text(at(1), TEXT_LIMITS["col2"]),
            "col3": clean_text(at(2), TEXT_LIMITS["col3"]),
            "amount": amount,
            "trxn_code": clean_text(at(positions["trxn_code"]), TEXT_LIMITS["trxn_code"]).upper(),
            "merchant": clean_text(at(positions["merchant"]), TEXT_LIMITS["merchant"]),
            "respond_code": clean_text(at(positions["respond_code"]), TEXT_LIMITS["respond_code"]),
        })
    return records, stats


# ---------------------------------------------------------------- duplicate rules

def find_duplicate_ids(rows):
    """Apply the OT/OS/OR/OM rules. `rows`: dicts with id, amount, trxn_code, merchant.

    Rows are grouped by (amount, merchant); rows without an amount are ignored.
      * OT and OS in the same group  -> every OS is flagged (the OT is kept)
      * OR or OM in the same group   -> every OT and OS is flagged
      * OT duplicates with no OS/OR/OM are kept
    """
    groups = defaultdict(lambda: defaultdict(list))
    for row in rows:
        if row["amount"] is None:
            continue
        key = (row["amount"], (row["merchant"] or "").strip())
        groups[key][(row["trxn_code"] or "").strip().upper()].append(row["id"])

    flagged = set()
    for codes in groups.values():
        if "OT" in codes and "OS" in codes:
            flagged.update(codes["OS"])
        if "OR" in codes or "OM" in codes:
            flagged.update(codes.get("OT", []))
            flagged.update(codes.get("OS", []))
    return flagged
